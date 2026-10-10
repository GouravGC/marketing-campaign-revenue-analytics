
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "processed"
REPORTS = ROOT / "reports" / "generated"
EXCEL = ROOT / "excel"
POWERBI = ROOT / "powerbi"

for folder in (REPORTS, EXCEL, POWERBI):
    folder.mkdir(parents=True, exist_ok=True)

files = {
    "events": DATA / "stg_ga4_events.csv",
    "daily": DATA / "mart_daily_event_performance.csv",
    "traffic": DATA / "mart_traffic_source_performance.csv",
    "summary": DATA / "mart_event_summary.csv",
}

for name, path in files.items():
    if not path.exists():
        raise FileNotFoundError(
            f"Missing {path}. Run python -m src.pipeline first."
        )

events = pd.read_csv(files["events"])
daily = pd.read_csv(files["daily"])
traffic = pd.read_csv(files["traffic"])
summary = pd.read_csv(files["summary"])

quality_path = REPORTS / "data_quality_results.csv"
quality = pd.read_csv(quality_path) if quality_path.exists() else pd.DataFrame()

duplicates_path = REPORTS / "duplicate_profile.txt"
duplicate_note = (
    duplicates_path.read_text(encoding="utf-8")
    if duplicates_path.exists() else "Duplicate profile not available."
)

# Power BI source files
exports = {
    "fact_ga4_events.csv": events,
    "mart_daily_performance.csv": daily,
    "mart_traffic_source_performance.csv": traffic,
    "mart_event_summary.csv": summary,
}
for filename, frame in exports.items():
    frame.to_csv(POWERBI / filename, index=False)

# Excel reconciliation workbook
workbook = EXCEL / "GA4_Analytics_Reconciliation.xlsx"
with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
    overview = pd.DataFrame({
        "Metric": [
            "Extract rows",
            "Distinct users",
            "Event types",
            "Purchase events",
            "Purchase revenue (USD)",
            "Exact duplicate rows in raw extract",
        ],
        "Value": [
            len(events),
            events["user_pseudo_id"].nunique(),
            events["event_name"].nunique(),
            int((events["event_name"] == "purchase").sum()),
            round(events["revenue_for_analytics_usd"].sum(), 2),
            int(pd.read_csv(ROOT / "data/raw/csv/ga4_ecommerce_events.csv")
                .duplicated().sum()),
        ],
    })
    overview.to_excel(writer, sheet_name="Overview", index=False)
    daily.to_excel(writer, sheet_name="Daily Performance", index=False)
    traffic.to_excel(writer, sheet_name="Traffic Sources", index=False)
    summary.to_excel(writer, sheet_name="Event Summary", index=False)
    quality.to_excel(writer, sheet_name="Data Quality", index=False)
    pd.DataFrame({"Notes": duplicate_note.splitlines()}).to_excel(
        writer, sheet_name="Duplicate Notes", index=False
    )

    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            letter = column[0].column_letter
            width = max(
                (len(str(cell.value)) for cell in column if cell.value is not None),
                default=10,
            )
            ws.column_dimensions[letter].width = min(width + 2, 38)

    ws = writer.book["Event Summary"]
    if ws.max_row > 1:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Events by Event Type"
        chart.y_axis.title = "Event"
        chart.x_axis.title = "Count"
        chart.add_data(Reference(ws, min_col=2, max_col=2,
                                 min_row=1, max_row=ws.max_row),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1,
                                       min_row=2, max_row=ws.max_row))
        ws.add_chart(chart, "E2")

print("REPORTING BUILD COMPLETE")
print(f"Excel workbook: {workbook}")
print(f"Power BI CSV exports: {POWERBI}")
print(f"Data quality report: {quality_path}")
